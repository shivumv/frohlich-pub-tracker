
import io
import re
import time
import string
import tempfile
import os
import requests
import pandas as pd
import streamlit as st
from bs4 import BeautifulSoup
from rapidfuzz import fuzz
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from xml.etree import ElementTree as ET

st.set_page_config(
    page_title="Frohlich Lab · Publication Tracker",
    page_icon="🧠",
    layout="wide",
)

PUBMED_QUERY = "Flavio Frohlich[Author]"
WEBSITE_URL = "https://www.frohlichlab.org/publications.html"
FUZZY_THRESHOLD = 90

NCBI_EMAIL = "lab_tool@frohlichlab.org"
NCBI_TOOL = "FrohlichLabCompare"

ESEARCH_URL = "https://eutils.ncbi.nlm.nih.gov/entrez/eutils/esearch.fcgi"
EFETCH_URL = "https://eutils.ncbi.nlm.nih.gov/entrez/eutils/efetch.fcgi"

COLUMNS = ["DOI", "PMID", "Title", "Year", "In PubMed", "On Website", "Match Status", "Notes"]

STATUS_COLORS = {
    "Found in both": "#C6EFCE",
    "Missing from website": "#FFC7CE",
    "Missing from PubMed": "#BDD7EE",
    "Metadata mismatch": "#FFEB9C",
}

FILL_MAP = {
    "Found in both": PatternFill("solid", fgColor="C6EFCE"),
    "Missing from website": PatternFill("solid", fgColor="FFC7CE"),
    "Missing from PubMed": PatternFill("solid", fgColor="BDD7EE"),
    "Metadata mismatch": PatternFill("solid", fgColor="FFEB9C"),
}

FILL_HEADER = PatternFill("solid", fgColor="2F75B6")


def normalize_title(t):
    if not t:
        return ""
    t = t.lower().translate(str.maketrans("", "", string.punctuation))
    return re.sub(r"\s+", " ", t).strip()


def normalize_doi(doi):
    if not doi:
        return ""
    doi = re.sub(r"^https?://(dx\.)?doi\.org/", "", doi, flags=re.IGNORECASE)
    return doi.lower().rstrip(".,)")


@st.cache_data(ttl=3600, show_spinner=False)
def fetch_pubmed(query):
    params = dict(db="pubmed", term=query, retmax=10000, retmode="json",
                  email=NCBI_EMAIL, tool=NCBI_TOOL)
    r = requests.get(ESEARCH_URL, params=params, timeout=30)
    r.raise_for_status()
    pmids = r.json()["esearchresult"]["idlist"]

    records = []
    for i in range(0, len(pmids), 200):
        batch = pmids[i:i + 200]
        p2 = dict(db="pubmed", id=",".join(batch), retmode="xml",
                  rettype="abstract", email=NCBI_EMAIL, tool=NCBI_TOOL)
        r2 = requests.get(EFETCH_URL, params=p2, timeout=60)
        r2.raise_for_status()
        records.extend(_parse_xml(r2.text))
        time.sleep(0.35)
    return records


def _parse_xml(xml_text):
    root = ET.fromstring(xml_text)
    out = []
    for article in root.iter("PubmedArticle"):
        pmid = getattr(article.find(".//PMID"), "text", "").strip()
        title = "".join((article.find(".//ArticleTitle") or ET.Element("x")).itertext()).strip()

        year = ""
        pd = article.find(".//PubDate")
        if pd is not None:
            y = pd.find("Year")
            if y is not None and y.text:
                year = y.text

        doi = ""
        for el in article.findall(".//ArticleId"):
            if el.get("IdType") == "doi":
                doi = normalize_doi(el.text or "")

        out.append({"PMID": pmid, "Title": title, "Year": year, "DOI": doi})
    return out


@st.cache_data(ttl=3600, show_spinner=False)
def scrape_website(url):
    r = requests.get(url, headers={"User-Agent": "Mozilla/5.0"}, timeout=30)
    soup = BeautifulSoup(r.text, "lxml")

    # get full text of page instead of per-tag parsing
    full_text = soup.get_text("\n")

    records = []
    current_year = None

    # split by lines
    lines = [l.strip() for l in full_text.split("\n") if l.strip()]

    buffer = []

    for line in lines:

        # detect year
        if re.match(r"^20\d{2}$", line):
            current_year = line
            continue

        # skip until we hit first year
        if not current_year:
            continue

        buffer.append(line)

        # heuristic: publication usually ends with period
        if line.endswith(".") and len(" ".join(buffer)) > 80:

            text = " ".join(buffer)
            buffer = []

            title = _extract_title(text)
            doi = _extract_doi_from_text(text)

            if title and len(text) > 80:
                records.append({
                    "Title": title,
                    "Year": current_year,
                    "DOI": normalize_doi(doi),
                    "Raw": text
                })

    return records

def _extract_doi_from_text(text):
    m = re.search(r"https?://(?:dx\.)?doi\.org/\S+", text)
    if m:
        return m.group(0)
    m2 = re.search(r"\b10\.\d{4,9}/\S+", text)
    return m2.group(0) if m2 else ""


def _extract_title(raw):
    parts = raw.split(". ")
    for p in parts:
        if 20 < len(p) < 300:
            return p
    return raw[:200]


def match_publications(pubmed, website):
    web_by_doi = {w["DOI"]: w for w in website if w.get("DOI")}
    web_by_title = {normalize_title(w["Title"]): w for w in website}

    rows = []

    for pm in pubmed:
        pm_title = normalize_title(pm["Title"])
        pm_doi = pm.get("DOI", "")

        match = None

        if pm_doi and pm_doi in web_by_doi:
            match = web_by_doi[pm_doi]

        if not match:
            best, best_score = None, 0
            for t, w in web_by_title.items():
                s = fuzz.token_sort_ratio(pm_title, t)
                if s > best_score:
                    best, best_score = w, s
            if best_score >= FUZZY_THRESHOLD:
                match = best

        if match:
            rows.append({**pm, "In PubMed": "Yes", "On Website": "Yes", "Match Status": "Found in both"})
        else:
            rows.append({**pm, "In PubMed": "Yes", "On Website": "No", "Match Status": "Missing from website"})

    return rows


st.title("Publication Tracker")

if st.button("Run"):
    pubmed = fetch_pubmed(PUBMED_QUERY)
    website = scrape_website(WEBSITE_URL)

    st.write(len(pubmed), "PubMed records")
    st.write(len(website), "Website records")

    rows = match_publications(pubmed, website)
    st.dataframe(pd.DataFrame(rows))
