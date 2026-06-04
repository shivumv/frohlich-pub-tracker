import io
import re
import time
import string
import tempfile
import os
import requests
import pandas as pd
import streamlit as st
from bs4 import BeautifulSoup
from rapidfuzz import fuzz
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from xml.etree import ElementTree as ET

st.set_page_config(
    page_title="Frohlich Lab · Publication Tracker",
    page_icon="🧠",
    layout="wide",
)

PUBMED_QUERY    = "Flavio Frohlich[Author]"
WEBSITE_URL     = "https://www.frohlichlab.org/publications.html"
FUZZY_THRESHOLD = 90
NCBI_EMAIL      = "lab_tool@frohlichlab.org"
NCBI_TOOL       = "FrohlichLabCompare"
ESEARCH_URL     = "https://eutils.ncbi.nlm.nih.gov/entrez/eutils/esearch.fcgi"
EFETCH_URL      = "https://eutils.ncbi.nlm.nih.gov/entrez/eutils/efetch.fcgi"

COLUMNS = ["DOI", "PMID", "Title", "Year", "In PubMed", "On Website", "Match Status", "Notes"]

STATUS_COLORS = {
    "Found in both":        "#C6EFCE",
    "Missing from website": "#FFC7CE",
    "Missing from PubMed":  "#BDD7EE",
    "Metadata mismatch":    "#FFEB9C",
}

FILL_MAP = {
    "Found in both":        PatternFill("solid", fgColor="C6EFCE"),
    "Missing from website": PatternFill("solid", fgColor="FFC7CE"),
    "Missing from PubMed":  PatternFill("solid", fgColor="BDD7EE"),
    "Metadata mismatch":    PatternFill("solid", fgColor="FFEB9C"),
}
FILL_HEADER = PatternFill("solid", fgColor="2F75B6")
COL_WIDTHS  = {"A": 40, "B": 12, "C": 70, "D": 8, "E": 12, "F": 12, "G": 24, "H": 45}


def normalize_title(t):
    if not t:
        return ""
    t = t.lower().translate(str.maketrans("", "", string.punctuation))
    return re.sub(r"\s+", " ", t).strip()


def normalize_doi(doi):
    if not doi:
        return ""
    doi = doi.strip()
    doi = re.sub(r"^https?://(dx\.)?doi\.org/", "", doi, flags=re.IGNORECASE)
    return doi.lower().rstrip(".,)")


@st.cache_data(ttl=3600, show_spinner=False)
def fetch_pubmed(query):
    params = dict(db="pubmed", term=query, retmax=10000, retmode="json",
                  email=NCBI_EMAIL, tool=NCBI_TOOL)
    r = requests.get(ESEARCH_URL, params=params, timeout=30)
    r.raise_for_status()
    pmids = r.json()["esearchresult"]["idlist"]
    records = []
    for i in range(0, len(pmids), 200):
        batch = pmids[i:i + 200]
        p2 = dict(db="pubmed", id=",".join(batch), retmode="xml",
                  rettype="abstract", email=NCBI_EMAIL, tool=NCBI_TOOL)
        r2 = requests.get(EFETCH_URL, params=p2, timeout=60)
        r2.raise_for_status()
        records.extend(_parse_xml(r2.text))
        time.sleep(0.35)
    return records


def _parse_xml(xml_text):
    root = ET.fromstring(xml_text)
    out = []
    for article in root.iter("PubmedArticle"):
        pmid  = getattr(article.find(".//PMID"), "text", "").strip()
        title = "".join((article.find(".//ArticleTitle") or ET.Element("x")).itertext()).strip()
        year  = ""
        pd_el = article.find(".//PubDate")
        if pd_el is not None:
            yr = pd_el.find("Year")
            year = yr.text.strip() if yr is not None else ""
            if not year:
                md = pd_el.find("MedlineDate")
                if md is not None:
                    m = re.search(r"\d{4}", md.text or "")
                    if m:
                        year = m.group()
        doi = ""
        for id_el in article.findall(".//ArticleId"):
            if id_el.get("IdType") == "doi":
                doi = normalize_doi(id_el.text or "")
                break
        authors = []
        for a in article.findall(".//Author"):
            ln = getattr(a.find("LastName"), "text", "") or ""
            fn = getattr(a.find("ForeName"), "text", "") or ""
            if ln:
                authors.append(f"{ln} {fn}".strip())
        out.append({"PMID": pmid, "Title": title, "Year": year,
                    "Authors": "; ".join(authors), "DOI": doi})
    return out


@st.cache_data(ttl=3600, show_spinner=False)
def scrape_website(url):
    headers = {"User-Agent": "FrohlichLabCompare/1.0"}
    r = requests.get(url, headers=headers, timeout=30)
    r.raise_for_status()
    soup = BeautifulSoup(r.text, "lxml")
    records, current_year = [], ""
    for tag in soup.find_all(["h2", "p"]):
        if tag.name == "h2":
            text = tag.get_text(strip=True)
            if re.match(r"^\d{4}$", text):
                current_year = text
        elif tag.name == "p" and current_year:
            raw = tag.get_text(separator=" ", strip=True)
            if len(raw) < 40:
                continue
            doi   = _extract_doi(tag)
            title = _extract_title(raw)
            if title:
                records.append({"Title": title, "Year": current_year,
                                 "DOI": normalize_doi(doi), "Raw": raw})
    return records


def _extract_doi(tag):
    for a in tag.find_all("a", href=True):
        if "doi.org" in a["href"]:
            return a["href"]
    text = tag.get_text(" ")
    m = re.search(r"https?://(?:dx\.)?doi\.org/(\S+?)(?:\s|$|\[)", text)
    if m:
        return m.group(0)
    m2 = re.search(r"\b(10\.\d{4,}/\S+?)(?:\s|$|\[|,|\))", text)
    return m2.group(1) if m2 else ""


def _extract_title(raw):
    m = re.search(r"\*(.+?)\*", raw)
    if m:
        return m.group(1).strip()
    parts = re.split(r"\.\s+", raw, maxsplit=3)
    for part in parts[1:]:
        c = part.strip()
        if 15 < len(c) < 300:
            return c
    return raw[:200].strip()


def _parse_plain_text(text):
    records = []
    current_year = ""
    for para in re.split(r"\n{2,}", text):
        para = para.strip()
        if not para:
            continue
        if re.match(r"^\d{4}$", para):
            current_year = para
            continue
        doi = ""
        m = re.search(r"https?://(?:dx\.)?doi\.org/(\S+?)(?:\s|$)", para)
        if m:
            doi = normalize_doi(m.group(0))
        else:
            m2 = re.search(r"\b(10\.\d{4,}/\S+?)(?:\s|[.,)]|$)", para)
            if m2:
                doi = normalize_doi(m2.group(1))
        title = _extract_title(para)
        if title and len(para) > 40:
            if not current_year:
                ym = re.search(r"\b(20\d{2})\b", para)
                if ym:
                    current_year = ym.group(1)
            records.append({"Title": title, "Year": current_year,
                             "DOI": doi, "Raw": para})
    return records


def load_pubmed_csv(uploaded_file):
    df = pd.read_csv(uploaded_file, dtype=str).fillna("")
    df.columns = [c.strip().lower() for c in df.columns]

    def pick(candidates):
        for c in candidates:
            if c in df.columns:
                return c
        return None

    t_col = pick(["title"])
    if t_col is None:
        raise ValueError("CSV has no 'Title' column.")
    records = []
    for _, row in df.iterrows():
        records.append({
            "Title":   row.get(t_col, ""),
            "PMID":    row.get(pick(["pmid", "pubmed id"]) or "", ""),
            "DOI":     normalize_doi(row.get(pick(["doi"]) or "", "")),
            "Year":    row.get(pick(["publication year", "year", "pub year"]) or "", ""),
            "Authors": row.get(pick(["authors", "author"]) or "", ""),
        })
    return records


def match_publications(pubmed, website):
    web_by_doi   = {r["DOI"]: r for r in website if r.get("DOI")}
    web_by_title = {normalize_title(r["Title"]): r for r in website}
    matched_web  = set()
    rows = []

    for pm in pubmed:
        pm_doi   = pm.get("DOI", "")
        pm_title = normalize_title(pm.get("Title", ""))
        web_match, match_type = None, ""

        if pm_doi and pm_doi in web_by_doi:
            web_match, match_type = web_by_doi[pm_doi], "DOI match"

        if web_match is None:
            best_score, best_entry = 0, None
            for wt, wr in web_by_title.items():
                score = fuzz.token_sort_ratio(pm_title, wt)
                if score > best_score:
                    best_score, best_entry = score, wr
            if best_score >= FUZZY_THRESHOLD:
                web_match  = best_entry
                match_type = f"Fuzzy title ({best_score:.0f}%)"

        if web_match is not None:
            matched_web.add(id(web_match))
            yr_pm  = str(pm.get("Year", "")).strip()
            yr_web = str(web_match.get("Year", "")).strip()
            if yr_pm and yr_web and yr_pm != yr_web:
                status = "Metadata mismatch"
                notes  = f"Year — PubMed: {yr_pm}, Website: {yr_web}"
            else:
                status, notes = "Found in both", match_type
            rows.append({"DOI": pm_doi or web_match.get("DOI", ""),
                         "PMID": pm.get("PMID", ""), "Title": pm.get("Title", ""),
                         "Year": yr_pm or yr_web,
                         "In PubMed": "Yes", "On Website": "Yes",
                         "Match Status": status, "Notes": notes,
                         "_status": status})
        else:
            rows.append({"DOI": pm_doi, "PMID": pm.get("PMID", ""),
                         "Title": pm.get("Title", ""), "Year": pm.get("Year", ""),
                         "In PubMed": "Yes", "On Website": "No",
                         "Match Status": "Missing from website", "Notes": "",
                         "_status": "Missing from website"})

    for wr in website:
        if id(wr) not in matched_web:
            rows.append({"DOI": wr.get("DOI", ""), "PMID": "",
                         "Title": wr.get("Title", ""), "Year": wr.get("Year", ""),
                         "In PubMed": "No", "On Website": "Yes",
                         "Match Status": "Missing from PubMed", "Notes": "",
                         "_status": "Missing from PubMed"})
    return rows


def build_excel(rows, pubmed, website):
    df = pd.DataFrame(rows)[COLUMNS]
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Comparison", index=False)
        _write_summary(writer, rows, pubmed, website)
    buf.seek(0)
    wb = load_workbook(buf)
    _format_comparison(wb["Comparison"], rows)
    _format_summary(wb["Summary"])
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


def _format_comparison(ws, rows):
    for cell in ws[1]:
        cell.fill      = FILL_HEADER
        cell.font      = Font(bold=True, color="FFFFFF", name="Arial", size=10)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.row_dimensions[1].height = 28
    for idx, rd in enumerate(rows, 2):
        fill = FILL_MAP.get(rd.get("_status", ""))
        for col in range(1, len(COLUMNS) + 1):
            cell = ws.cell(row=idx, column=col)
            cell.font      = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if fill:
                cell.fill = fill
    for letter, width in COL_WIDTHS.items():
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A2"
    lr = ws.max_row + 2
    ws.cell(row=lr, column=1, value="Legend").font = Font(bold=True, name="Arial")
    for i, (label, fill, desc) in enumerate([
        ("Green",  FILL_MAP["Found in both"],        "Found in both sources"),
        ("Red",    FILL_MAP["Missing from website"],  "In PubMed, missing from website"),
        ("Blue",   FILL_MAP["Missing from PubMed"],   "On website, missing from PubMed"),
        ("Yellow", FILL_MAP["Metadata mismatch"],     "Metadata mismatch (e.g. year)"),
    ], 1):
        r = lr + i
        c = ws.cell(row=r, column=1, value=label)
        c.fill = fill
        c.font = Font(name="Arial", size=10)
        ws.cell(row=r, column=2, value=desc).font = Font(name="Arial", size=10)


def _write_summary(writer, rows, pubmed, website):
    n_both   = sum(1 for r in rows if r["In PubMed"] == "Yes" and r["On Website"] == "Yes")
    n_doi    = sum(1 for r in rows if "DOI match"    in r.get("Notes", ""))
    n_fuzzy  = sum(1 for r in rows if "Fuzzy title"  in r.get("Notes", ""))
    n_miss_w = sum(1 for r in rows if r["Match Status"] == "Missing from website")
    n_miss_p = sum(1 for r in rows if r["Match Status"] == "Missing from PubMed")
    n_mis    = sum(1 for r in rows if r["Match Status"] == "Metadata mismatch")
    pd.DataFrame({
        "Metric": ["Total in PubMed", "Total on website", "Found in both",
                   "Exact DOI matches", "Fuzzy title matches", "Metadata mismatches",
                   "Missing from website", "Missing from PubMed"],
        "Count":  [len(pubmed), len(website), n_both, n_doi, n_fuzzy,
                   n_mis, n_miss_w, n_miss_p],
    }).to_excel(writer, sheet_name="Summary", index=False)


def _format_summary(ws):
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 12
    for cell in ws[1]:
        cell.fill      = FILL_HEADER
        cell.font      = Font(bold=True, color="FFFFFF", name="Arial", size=11)
        cell.alignment = Alignment(horizontal="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font      = Font(name="Arial", size=11)
            cell.alignment = Alignment(horizontal="left" if cell.column == 1 else "center")
    ws.freeze_panes = "A2"


st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=DM+Serif+Display&family=DM+Sans:wght@400;500;600&display=swap');

html, body, [class*="css"] { font-family: 'DM Sans', sans-serif; }
h1, h2, h3 { font-family: 'DM Serif Display', serif !important; }

.hero {
    background: linear-gradient(135deg, #0f2744 0%, #1a3a5c 60%, #0d5c8f 100%);
    border-radius: 16px;
    padding: 2.5rem 2.5rem 2rem;
    margin-bottom: 2rem;
    color: white;
}
.hero h1 { color: white !important; font-size: 2.2rem; margin: 0 0 .4rem; }
.hero p  { color: #a8c8e8; margin: 0; font-size: 1.05rem; }

.stat-card {
    background: white;
    border: 1px solid #e2e8f0;
    border-radius: 12px;
    padding: 1.2rem 1rem;
    text-align: center;
    box-shadow: 0 1px 4px rgba(0,0,0,.06);
}
.stat-card .number { font-size: 2rem; font-weight: 700; line-height: 1; }
.stat-card .label  { font-size: .78rem; color: #64748b; margin-top: .3rem; text-transform: uppercase; letter-spacing: .04em; }
.green-num  { color: #16a34a; }
.red-num    { color: #dc2626; }
.blue-num   { color: #2563eb; }
.yellow-num { color: #ca8a04; }
.total-num  { color: #0f2744; }

.legend-row { display:flex; align-items:center; gap:.6rem; margin:.3rem 0; font-size:.9rem; }
.legend-dot { width:14px; height:14px; border-radius:3px; flex-shrink:0; }
</style>
""", unsafe_allow_html=True)

st.markdown("""
<div class="hero">
  <h1>🧠 Publication Tracker</h1>
  <p>Frohlich Lab · Automated comparison of PubMed vs. lab website</p>
</div>
""", unsafe_allow_html=True)

with st.sidebar:
    st.markdown("### ⚙️ Settings")
    pubmed_query = st.text_input("PubMed search query", value=PUBMED_QUERY,
                                 help="Any valid PubMed search string")
    fuzzy_thresh = st.slider("Fuzzy match threshold", 70, 100, FUZZY_THRESHOLD,
                             help="Minimum title similarity score (0–100)")
    st.markdown("---")
    st.markdown("### 📂 Use local files (optional)")
    st.caption("Skip live fetching by uploading exports instead.")
    csv_file = st.file_uploader("PubMed CSV export", type=["csv"])
    txt_file = st.file_uploader("Lab website PDF or text", type=["pdf", "txt"])
    st.markdown("---")
    st.markdown("""
    <div style='font-size:.82rem;color:#64748b;'>
    <b>Legend</b><br>
    <div class='legend-row'><div class='legend-dot' style='background:#C6EFCE'></div> Found in both</div>
    <div class='legend-row'><div class='legend-dot' style='background:#FFC7CE'></div> PubMed only</div>
    <div class='legend-row'><div class='legend-dot' style='background:#BDD7EE'></div> Website only</div>
    <div class='legend-row'><div class='legend-dot' style='background:#FFEB9C'></div> Metadata mismatch</div>
    </div>
    """, unsafe_allow_html=True)

run = st.button("▶  Run Comparison", type="primary", use_container_width=True)

if run:
    errors = []

    with st.status("Fetching PubMed records …", expanded=True):
        try:
            if csv_file:
                st.write("Loading PubMed CSV …")
                pubmed = load_pubmed_csv(csv_file)
            else:
                st.write(f"Querying PubMed: `{pubmed_query}` …")
                pubmed = fetch_pubmed(pubmed_query)
            st.write(f"✅ {len(pubmed)} PubMed records loaded")
        except Exception as e:
            errors.append(f"PubMed error: {e}")
            pubmed = []
            st.write(f"❌ {e}")

    with st.status("Scraping lab website …", expanded=True):
        try:
            if txt_file:
                suffix = txt_file.name.split(".")[-1].lower()
                if suffix == "pdf":
                    import pdfplumber
                    with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp:
                        tmp.write(txt_file.read())
                        tmp_path = tmp.name
                    pages = []
                    with pdfplumber.open(tmp_path) as pdf:
                        for page in pdf.pages:
                            t = page.extract_text()
                            if t:
                                pages.append(t)
                    os.unlink(tmp_path)
                    website = _parse_plain_text("\n".join(pages))
                else:
                    text = txt_file.read().decode("utf-8", errors="replace")
                    website = _parse_plain_text(text)
            else:
                st.write(f"Scraping `{WEBSITE_URL}` …")
                website = scrape_website(WEBSITE_URL)
            st.write(f"✅ {len(website)} website entries loaded")
        except Exception as e:
            errors.append(f"Website error: {e}")
            website = []
            st.write(f"❌ {e}")

    if errors:
        for err in errors:
            st.error(err)
        st.stop()

    with st.spinner("Matching publications …"):
        rows = match_publications(pubmed, website)

    n_green  = sum(1 for r in rows if r["_status"] == "Found in both")
    n_yellow = sum(1 for r in rows if r["_status"] == "Metadata mismatch")
    n_red    = sum(1 for r in rows if r["_status"] == "Missing from website")
    n_blue   = sum(1 for r in rows if r["_status"] == "Missing from PubMed")

    c1, c2, c3, c4, c5, c6 = st.columns(6)
    for col, num, label, css in [
        (c1, len(pubmed),  "PubMed total",        "total-num"),
        (c2, len(website), "Website total",        "total-num"),
        (c3, n_green,      "Found in both",        "green-num"),
        (c4, n_red,        "Missing from website", "red-num"),
        (c5, n_blue,       "Missing from PubMed",  "blue-num"),
        (c6, n_yellow,     "Metadata mismatch",    "yellow-num"),
    ]:
        with col:
            st.markdown(f"""
            <div class="stat-card">
              <div class="number {css}">{num}</div>
              <div class="label">{label}</div>
            </div>""", unsafe_allow_html=True)

    st.markdown("<br>", unsafe_allow_html=True)

    tab_all, tab_both, tab_red, tab_blue, tab_yellow = st.tabs([
        f"All ({len(rows)})",
        f"✅ Both ({n_green})",
        f"🔴 Missing from website ({n_red})",
        f"🔵 Missing from PubMed ({n_blue})",
        f"🟡 Mismatch ({n_yellow})",
    ])

    def render_table(filtered_rows):
        if not filtered_rows:
            st.info("No records in this category.")
            return
        df = pd.DataFrame(filtered_rows)[COLUMNS]

        def color_row(row):
            color = STATUS_COLORS.get(row["Match Status"], "")
            return [f"background-color: {color}" if color else ""] * len(row)

        styled = df.style.apply(color_row, axis=1)
        st.dataframe(styled, use_container_width=True, height=420,
                     column_config={
                         "DOI":   st.column_config.LinkColumn("DOI", display_text="🔗 Link"),
                         "Title": st.column_config.TextColumn("Title", width="large"),
                     })

    with tab_all:    render_table(rows)
    with tab_both:   render_table([r for r in rows if r["_status"] == "Found in both"])
    with tab_red:    render_table([r for r in rows if r["_status"] == "Missing from website"])
    with tab_blue:   render_table([r for r in rows if r["_status"] == "Missing from PubMed"])
    with tab_yellow: render_table([r for r in rows if r["_status"] == "Metadata mismatch"])

    st.markdown("---")
    excel_bytes = build_excel(rows, pubmed, website)
    st.download_button(
        label="⬇️  Download Excel Report",
        data=excel_bytes,
        file_name="frohlich_publications_report.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
        type="primary",
    )

else:
    st.info("👆 Click **Run Comparison** to fetch PubMed and scrape the lab website automatically.")
    st.markdown("""
    **What this tool does:**
    - Queries PubMed for all Flavio Frohlich publications via the NCBI API
    - Scrapes the Frohlich Lab website publications page
    - Matches them by DOI (exact) or title similarity (fuzzy, ≥90%)
    - Highlights discrepancies in a colour-coded table
    - Exports a formatted Excel report with a Summary sheet
    """)
